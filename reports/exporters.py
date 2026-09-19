from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def export_documents_to_excel(queryset) -> BytesIO:
    """Выгрузить документы в Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Документы"

    headers = [
        "ID",
        "Название",
        "Тип",
        "Статус",
        "Владелец",
        "Номер",
        "Дата",
        "Контрагент",
        "Сумма",
        "Валюта",
        "Создан",
        "Обновлён",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="4F81BD")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for doc in queryset.select_related("owner"):
        fields = (doc.extracted_data or {}).get("fields", {})
        ws.append(
            [
                doc.id,
                doc.title,
                doc.get_doc_type_display(),
                doc.get_status_display(),
                doc.owner.username,
                fields.get("doc_number") or "",
                fields.get("doc_date") or "",
                fields.get("counterparty") or "",
                fields.get("amount") or "",
                fields.get("currency") or "",
                doc.created_at.strftime("%Y-%m-%d %H:%M"),
                doc.updated_at.strftime("%Y-%m-%d %H:%M"),
            ]
        )

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
